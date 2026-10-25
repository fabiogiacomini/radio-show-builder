import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from radio_show_builder import providers
from radio_show_builder.model import ValidationError, entries, new_project, settings, update_project
from radio_show_builder.store import Store
from radio_show_builder.documents import export_docx
from docx import Document

class CoreTests(unittest.TestCase):
    def test_presets_and_partial_updates_preserve_editor_choices(self):
        podcast = new_project()
        self.assertEqual((podcast["preset"], podcast["music_percent"], podcast["settings"]["duration_mode"]),
                         ("podcast", 0, "natural"))
        self.assertEqual(podcast["settings"]["target_seconds"], 1800)
        radio = new_project(preset="radio")
        self.assertEqual((radio["music_percent"], radio["settings"]["duration_mode"],
                          radio["settings"]["target_seconds"]), (83, "exact", 5400))
        updated = update_project(podcast, {"settings": {"bitrate_kbps": 256}, "music_percent": 12})
        self.assertEqual(updated["settings"]["duration_mode"], "natural")
        self.assertEqual(updated["settings"]["target_seconds"], 1800)
        self.assertEqual(updated["music_percent"], 12)
        legacy = dict(radio)
        del legacy["preset"], legacy["music_percent"]
        self.assertEqual(update_project(legacy, {"title": "Still on air"})["settings"], radio["settings"])
        for value in (-1, 101, True, None, "bad", float("nan")):
            with self.subTest(value=value), self.assertRaises(ValidationError):
                update_project(podcast, {"music_percent": value})
        for value in ("invalid", [], None):
            with self.subTest(value=value), self.assertRaises(ValidationError):
                new_project(preset=value)

    def test_nonfinite_settings_and_invalid_trims_are_rejected(self):
        for value in (float("nan"), float("inf"), -1):
            with self.assertRaises(ValidationError):
                settings({"target_seconds": value})
        with self.assertRaises(ValidationError):
            entries([{"kind":"music", "trim_start":10, "trim_end":5}])
        with self.assertRaises(ValidationError):
            entries([{"kind":"command"}])

    def test_duplicate_ids_and_unknown_asset_rejected(self):
        entry = entries([{"kind":"speech"}])[0]
        with self.assertRaises(ValidationError):
            entries([entry, entry])
        entry["asset_id"] = "a"*32
        with self.assertRaises(ValidationError):
            entries([entry], {})

    def test_project_does_not_persist_credential_fields(self):
        p = new_project("A public example")
        p = update_project(p, {"api_key":"not-a-real-key", "provider":{"api_key":"not-a-real-key"}})
        self.assertNotIn("api_key", json.dumps(p))

    def test_store_roundtrip_and_path_validation(self):
        with tempfile.TemporaryDirectory() as tmp:
            store = Store(tmp)
            p = store.create("An example show")
            self.assertEqual(store.load(p["id"]), p)
            self.assertEqual(len(store.list()), 1)
            with self.assertRaises(ValidationError):
                store.load("../project")
            with self.assertRaises(ValidationError):
                store.load("b"*32)

    def test_word_contains_script_and_generic_metadata(self):
        p = new_project("A journey through sound")
        p["idea"] = "Explore how instruments shape a musical scene."
        p["entries"] = entries([
            {"kind":"speech","title":"Opening","text":"Welcome. Today we follow a sound.\n\nLet us begin with the music."},
            {"kind":"music","title":"Example Artist — Example Track","estimated_seconds":240},
            {"kind":"speech","title":"Closing","text":"Thank you for listening."},
        ])
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp)/"script.docx"
            export_docx(p, path)
            doc = Document(path)
            self.assertIn("Welcome. Today", "\n".join(x.text for x in doc.paragraphs))
            self.assertEqual(doc.core_properties.author, "Radio Show Builder")
            self.assertEqual(len(doc.tables[0].rows), 4)

class ProviderTests(unittest.TestCase):
    def draft(self):
        return {"title":"A sonic journey","verification_notes":"Verify release dates.",
                "entries":[{"kind":"speech","title":"Opening","text":"Hello."},
                           {"kind":"music","title":"An example track","estimated_seconds":240}]}

    def test_all_three_provider_formats(self):
        p = new_project()
        p["idea"] = "A short programme about synthesisers"
        for protocol in ("chat","responses","anthropic"):
            raw = json.dumps(self.draft())
            response = {
                "chat":{"choices":[{"message":{"content":raw},"finish_reason":"stop"}]},
                "responses":{"status":"completed","output":[{"type":"message","content":[{"type":"output_text","text":raw}]}]},
                "anthropic":{"stop_reason":"end_turn","content":[{"type":"text","text":raw}]},
            }[protocol]
            with self.subTest(protocol=protocol), patch.object(providers, "request_json", return_value=response) as call:
                result = providers.generate(p, {"provider":"custom","protocol":protocol,
                    "base_url":"https://example.com/v1","model":"example-model","api_key":"not-a-real-key"})
                self.assertEqual(len(result["entries"]), 2)
                url, headers, body = call.call_args.args
                self.assertTrue(url.endswith({"chat":"/chat/completions","responses":"/responses","anthropic":"/messages"}[protocol]))
                self.assertEqual(body["model"], "example-model")
                brief = json.loads(body["input"] if protocol == "responses" else body["messages"][-1]["content"])
                self.assertEqual((brief["format"], brief["music_percent"], brief["duration_mode"]),
                                 ("podcast", 0, "natural"))
                if protocol=="anthropic":
                    self.assertEqual(headers["x-api-key"], "not-a-real-key")
                else:
                    self.assertEqual(headers["Authorization"], "Bearer not-a-real-key")
                self.assertNotIn("api_key", result)

    def test_legacy_radio_brief_keeps_music_and_timing(self):
        p = new_project(preset="radio")
        del p["preset"], p["music_percent"]
        response = {"output_text": json.dumps(self.draft())}
        with patch.object(providers, "request_json", return_value=response) as call:
            providers.generate(p, {"provider": "openai", "model": "example-model"})
            brief = json.loads(call.call_args.args[2]["input"])
            self.assertEqual((brief["format"], brief["music_percent"], brief["duration_minutes"]),
                             ("radio", 83, 90))

    def test_endpoint_credentials_and_insecure_remote_urls_rejected(self):
        for url in ("http://remote.example/v1","https://key@example.com/v1","https://example.com/v1?key=abc","file:///etc/passwd"):
            with self.subTest(url=url), self.assertRaises(ValidationError):
                providers.config({"base_url":url})
        self.assertEqual(providers.config({"provider":"ollama"})["base_url"],"http://localhost:11434/v1")

    def test_invalid_and_truncated_drafts_rejected(self):
        for raw in ('not json','{"entries":[]}', '{"entries":[{"kind":"system"}]}'):
            with self.assertRaises(ValidationError):
                providers.parse_draft(raw)
        with self.assertRaises(ValidationError):
            providers.extract({"choices":[{"finish_reason":"length","message":{"content":"{}"}}]}, "chat")

    def test_model_cannot_assign_existing_audio(self):
        draft = self.draft()
        draft["entries"][0].update(asset_id="a"*32, trim_start=2, trim_end=3)
        result = providers.parse_draft(json.dumps(draft))
        self.assertIsNone(result["entries"][0]["asset_id"])
        self.assertEqual(result["entries"][0]["trim_start"],0)
